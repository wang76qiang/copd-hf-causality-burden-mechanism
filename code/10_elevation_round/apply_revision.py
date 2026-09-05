# -*- coding: utf-8 -*-
"""Apply JTM 升华 revision to the MedComm manuscript docx -> ECM1_Manuscript_JTM.docx"""
import sys, shutil
sys.path.insert(0, ".")
from revision_content import *
from docx import Document

SRC = "../ECM1_Manuscript_MedComm.docx"
DST = "../ECM1_Manuscript_JTM.docx"
shutil.copy(SRC, DST)
doc = Document(DST)
paras = doc.paragraphs

def set_text(i, new, expect_prefix=None):
    p = paras[i]
    if expect_prefix and not p.text.startswith(expect_prefix):
        raise SystemExit(f"paragraph {i} mismatch: {p.text[:60]!r} != {expect_prefix!r}")
    # keep first run formatting; clear others
    for r in p.runs[1:]:
        r.text = ""
    if p.runs:
        p.runs[0].text = new
    else:
        p.add_run(new)

def delete(i):
    p = paras[i]
    p._element.getparent().remove(p._element)

def insert_before(anchor_idx, items):
    """items: list of (style, text). Insert before paras[anchor_idx]."""
    anchor = paras[anchor_idx]
    for style, text in items:
        np = anchor.insert_paragraph_before(text)
        np.style = doc.styles[style]

# ---------- replacements ----------
set_text(0, TITLE, "Chronic obstructive")
set_text(36, ABSTRACT_BACKGROUND, "Background Heart failure")
set_text(37, ABSTRACT_METHODS, "Methods We performed")
set_text(38, ABSTRACT_RESULTS, "Findings Genetically predicted")
set_text(39, ABSTRACT_CONCLUSIONS, "Interpretation COPD is")
set_text(40, KEYWORDS, "Keywords Chronic")
set_text(45, "Background", "Introduction")
set_text(48, INTRO_PARA3, "We reasoned that these three gaps")
# methods data sources: append
set_text(51, paras[51].text + METHODS_DATASOURCES_ADD, "The study followed a pre-specified")
# transcriptomics methods: append
set_text(61, paras[61].text + METHODS_SC_ADD, "Bulk differential expression")
# replication results ending
old68 = paras[68].text
assert old68.endswith("return to its sources in the Discussion.")
set_text(68, old68.replace("We interpret this pattern as a genuine but context-dependent causal signal and return to its sources in the Discussion.",
         RESULTS_REPLICATION_ENDING.strip()), "We attempted replication")
# druggability append
set_text(84, paras[84].text + DRUGGABILITY_ADD, "Four orthogonal lines")
# discussion
set_text(86, DISC_PARA1, "This study delivers")
set_text(87, DISC_PARA2, "The replication paradox")
set_text(88, DISC_PARA3, "In relation to prior work")
set_text(89, paras[89].text + DISC_PARA4_ADD, "Our results support a two-stage")
set_text(91, DISC_LIMIT1, "Several limitations")
set_text(92, DISC_NEXT, "Genetically, HFpEF-stratified")
set_text(93, CONCLUSIONS, "In conclusion, COPD is")
# declarations
set_text(29, "Availability of data and materials", "Data Availability Statement")
set_text(30, DATA_AVAIL, "All source datasets")

# ---------- deletions (Research in context) ----------
for i in [41, 42, 43, 44]:
    delete(i)
paras = doc.paragraphs  # refresh

# ---------- insertions ----------
def find_para(prefix, style_sub=None):
    for i, p in enumerate(paras):
        if p.text.startswith(prefix):
            return i
    raise SystemExit(f"anchor not found: {prefix}")

# new Methods subsections before "Global burden analysis" heading
i = find_para("Global burden analysis")
insert_before(i, [("Heading 2", METH_H2_WINNERS), ("Normal", METH_P_WINNERS),
                  ("Heading 2", METH_H2_INTERORGAN), ("Normal", METH_P_INTERORGAN)])
paras = doc.paragraphs
# new Results subsections before "The effect is independent of smoking"
i = find_para("The effect is independent of smoking")
insert_before(i, [("Heading 2", RES_H2_WINNERS), ("Normal", RES_P_WINNERS),
                  ("Heading 2", RES_H2_EXTENDED), ("Normal", RES_P_EXTENDED),
                  ("Heading 2", RES_H2_HETEROG), ("Normal", RES_P_HETEROG)])
paras = doc.paragraphs
# interorgan results before "Druggability"
i = find_para("Druggability: tissue-confined")
insert_before(i, [("Heading 2", RES_H2_INTERORGAN), ("Normal", RES_P_INTERORGAN)])
paras = doc.paragraphs
# Conclusions heading before final conclusion paragraph
i = find_para("In conclusion, COPD is")
insert_before(i, [("Heading 1", "Conclusions")])
paras = doc.paragraphs
# Consent for publication after Ethics paragraph
i = find_para("This study analysed publicly available, de-identified summary statistics and public datasets. All contributing")
insert_before(i + 1, [("Heading 1", "Consent for publication"), ("Normal", "Not applicable.")])
paras = doc.paragraphs
# new references before "Tables" heading
i = find_para("Tables")
ref_style = paras[i - 1].style.name
insert_before(i, [(ref_style, r) for r in NEW_REFS])
paras = doc.paragraphs
# Figure 8 legend at end (after Figure 7 legend)
fig7 = find_para("Figure 7. Cellular mechanism")
style_leg = paras[fig7].style.name
newp = doc.add_paragraph(FIG8_LEGEND)
newp.style = doc.styles[style_leg]

doc.save(DST)
print("saved", DST)
print("total paragraphs:", len(doc.paragraphs))
